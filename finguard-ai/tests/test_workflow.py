import io
import json
import sys
import time
from pathlib import Path
import fitz
import openpyxl
import pytest
from docx import Document
from fastapi.testclient import TestClient

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from finance import verify, calculate, screen
import workflow
import server


def fig(field="current_assets", value="80"):
    return dict(
        field=field,
        value=value,
        period="2025",
        currency="USD",
        scale=1,
        source_id="page:1",
        quote="Current assets 80; Current liabilities 100",
    )


def pdf_bytes():
    with fitz.open() as doc:
        p = doc.new_page()
        p.insert_text((50, 70), fig()["quote"])
        return doc.tobytes()


def test_evidence():
    sources = [{"id": "page:1", "text": fig()["quote"]}]
    assert len(verify([fig()], sources)[0]) == 1
    for change in (
        {"value": "900"},
        {"source_id": "page:2"},
        {"quote": "fabricated 80"},
        {"value": "NaN"},
        {"scale": 2},
    ):
        f = fig()
        f.update(change)
        assert verify([f], sources)[1]


def test_ratios_units_missing_and_rules():
    a, b = fig(), fig("current_liabilities", "100")
    r, m = calculate([a, b])
    assert r["current_ratio"]["value"] == "0.8000"
    assert screen([a, b], r)[0]["code"] == "LIQUIDITY"
    assert len(m) == 3
    a["scale"] = 1000
    assert calculate([a, b])[0]["current_ratio"]["value"] == "800.0000"
    b["value"] = "0"
    assert "current_ratio" not in calculate([a, b])[0]
    b["value"] = "100"
    b["period"] = "2024"
    assert "current_ratio" not in calculate([a, b])[0]
    assert "current_ratio" not in calculate([a, a, b])[0]


def test_thresholds_and_cash():
    fs = [
        fig("current_assets", "100"),
        fig("current_liabilities", "100"),
        fig("total_assets", "100"),
        fig("total_liabilities", "70"),
        fig("revenue", "100"),
        fig("accounts_receivable", "35"),
        fig("net_profit", "10"),
        fig("operating_cash_flow", "-1"),
    ]
    assert [s["code"] for s in screen(fs, calculate(fs)[0])] == [
        "CASH_PROFIT_DIVERGENCE"
    ]
    fs[3]["value"] = "71"
    assert "LEVERAGE" in [s["code"] for s in screen(fs, calculate(fs)[0])]


def test_parsers():
    assert workflow.extract(pdf_bytes(), "annual.pdf")[0]["id"] == "page:1"
    d = Document()
    d.add_paragraph("Revenue 100")
    table = d.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Revenue"
    table.cell(0, 1).text = "100"
    stream = io.BytesIO()
    d.save(stream)
    assert any(
        s["id"] == "table:1:row:1"
        for s in workflow.extract(stream.getvalue(), "a.docx")
    )
    wb = openpyxl.Workbook()
    wb.active.append(["Revenue", 100])
    stream = io.BytesIO()
    wb.save(stream)
    assert "B1=100" in workflow.extract(stream.getvalue(), "a.xlsx")[0]["text"]
    with pytest.raises(ValueError):
        workflow.extract(b"bad", "a.pdf")
    with fitz.open() as blank:
        blank.new_page()
        with pytest.raises(ValueError):
            workflow.extract(blank.tobytes(), "scan.pdf")


@pytest.mark.asyncio
async def test_agent_tools_and_citation_filter(monkeypatch):
    calls = iter(
        [
            {
                "content": None,
                "tool_calls": [
                    {
                        "id": "call1",
                        "type": "function",
                        "function": {"name": "calculate_metrics", "arguments": "{}"},
                    }
                ],
            },
            {
                "content": json.dumps(
                    {
                        "observations": [
                            {"text": "Review liquidity", "source_ids": ["page:1"]},
                            {"text": "Invented evidence", "source_ids": ["page:99"]},
                            {"text": "Profit 99", "source_ids": ["page:1"]},
                        ]
                    }
                )
            },
        ]
    )

    async def fake(*a, **kw):
        return next(calls)

    async def event(m):
        pass

    monkeypatch.setattr(workflow, "completion", fake)
    r = await workflow.agent([fig(), fig("current_liabilities", "100")], "cash", event)
    assert len(r["observations"]) == 1 and r["discarded_observations"] == 2
    assert r["tool_trace"][0]["tool"] == "calculate_metrics"


@pytest.mark.asyncio
async def test_agent_limits(monkeypatch):
    async def event(m):
        pass

    async def unknown(*args):
        return {
            "tool_calls": [
                {"id": "x", "function": {"name": "shell", "arguments": "{}"}}
            ]
        }

    monkeypatch.setattr(workflow, "completion", unknown)
    with pytest.raises(ValueError, match="Unknown"):
        await workflow.agent([fig()], "", event)

    async def loop(*args):
        return {
            "tool_calls": [
                {"id": "x", "function": {"name": "screen_risks", "arguments": "{}"}}
            ]
        }

    monkeypatch.setattr(workflow, "completion", loop)
    with pytest.raises(ValueError, match="iteration"):
        await workflow.agent([fig()], "", event)


def poll(c, token, status):
    for _ in range(100):
        d = c.get("/api/tasks/" + token).json()
        if d["status"] == status:
            return d
        time.sleep(0.01)
    raise AssertionError(d)


def test_api_pdf_persistence_security(monkeypatch, tmp_path):
    monkeypatch.setenv("FINGUARD_DB", str(tmp_path / "tasks.db"))
    monkeypatch.setenv("LLM_API_KEY", "test-only")
    monkeypatch.delenv("FINGUARD_ACCESS_TOKEN", raising=False)

    async def completion(messages, tools=None):
        return (
            {"content": '{"observations":[]}'}
            if tools
            else {
                "content": json.dumps(
                    {"figures": [fig(), fig("current_liabilities", "100")]}
                )
            }
        )

    monkeypatch.setattr(workflow, "completion", completion)
    with TestClient(server.app) as c:
        assert c.get("/health").status_code == 200
        assert (
            c.post("/api/tasks", files={"file": ("a.pdf", pdf_bytes())}).status_code
            == 422
        )
        assert (
            c.post(
                "/api/tasks",
                data={"provider_consent": "true"},
                files={"file": ("a.exe", b"bad")},
            ).status_code
            == 415
        )
        response = c.post(
            "/api/tasks",
            data={"provider_consent": "true"},
            files={"file": ("a.pdf", pdf_bytes())},
        )
        assert response.status_code == 202
        token = response.json()["id"]
        task = poll(c, token, "needs_review")
        assert c.get("/api/tasks/" + token + "/pdf").status_code == 409
        bad = fig()
        bad["value"] = "999"
        assert (
            c.post(
                "/api/tasks/" + token + "/confirm", json={"figures": [bad]}
            ).status_code
            == 422
        )
        assert (
            c.post(
                "/api/tasks/" + token + "/confirm", json={"figures": task["figures"]}
            ).status_code
            == 202
        )
        done = poll(c, token, "completed")
        assert done["result"]["ratios"]["current_ratio"]["value"] == "0.8000"
        output = c.get("/api/tasks/" + token + "/pdf")
        assert output.content.startswith(b"%PDF")
        with fitz.open(stream=output.content, filetype="pdf") as doc:
            assert "FinGuard" in doc[0].get_text()
        assert (
            c.post(
                "/api/tasks/" + token + "/confirm", json={"figures": task["figures"]}
            ).status_code
            == 409
        )
    with TestClient(server.app) as c:
        assert c.get("/api/tasks/" + token).json()["status"] == "completed"
        monkeypatch.setenv("FINGUARD_ACCESS_TOKEN", "test-token")
        assert c.get("/api/tasks").status_code == 401
        assert (
            c.get(
                "/api/tasks", headers={"Authorization": "Bearer test-token"}
            ).status_code
            == 200
        )
        assert (
            c.delete(
                "/api/tasks/" + token, headers={"Authorization": "Bearer test-token"}
            ).status_code
            == 204
        )


def test_missing_key_and_failure(monkeypatch, tmp_path):
    monkeypatch.setenv("FINGUARD_DB", str(tmp_path / "tasks.db"))
    monkeypatch.delenv("LLM_API_KEY", raising=False)
    monkeypatch.delenv("FINGUARD_ACCESS_TOKEN", raising=False)
    with TestClient(server.app) as c:
        assert (
            c.post(
                "/api/tasks",
                data={"provider_consent": "true"},
                files={"file": ("a.pdf", pdf_bytes())},
            ).status_code
            == 503
        )
        monkeypatch.setenv("LLM_API_KEY", "test-only")

        async def fail(*a, **kw):
            raise RuntimeError("secret-provider-content")

        monkeypatch.setattr(workflow, "completion", fail)
        d = c.post(
            "/api/tasks",
            data={"provider_consent": "true"},
            files={"file": ("a.pdf", pdf_bytes())},
        ).json()
        assert "secret-provider-content" not in json.dumps(poll(c, d["id"], "failed"))
